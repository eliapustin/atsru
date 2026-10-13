"""Собирает лист приёмки сайта для заказчика: docs/site-acceptance-checklist.xlsx.

Запуск: python docs/build-acceptance-checklist.py [путь_к_xlsx]
Пункты проверки — в списке items; при изменении структуры сайта правьте его.
"""
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1] if len(sys.argv) > 1 else Path(__file__).with_name('site-acceptance-checklist.xlsx')
F = 'Arial'
font = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
hfont = Font(name=F, size=10, bold=True, color='FFFFFF')
title = Font(name=F, size=14, bold=True)
hfill = PatternFill('solid', fgColor='1F3864')
secfill = PatternFill('solid', fgColor='D9E1F2')
inputfill = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='A6A6A6')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')

STATUSES = ['Принято', 'Замечание', 'Критично', 'Не проверялось']

H = 'https://aviatechnosoft.ru/'
items = [
    ('Общее', [
        ('Все страницы', 'Состав сайта соответствует ТЗ', 'Есть все разделы: Главная, Продукты (5 страниц), Для образования, Как закупить, Документы, О компании, Контакты, Оборудование и компоненты, Новости'),
        ('Все страницы', 'Шапка и главное меню', 'Логотип ведёт на главную, все пункты меню открывают нужные страницы, текущий раздел выделен'),
        ('Все страницы', 'Подвал сайта', 'Контакты, ссылки на разделы и копирайт актуальны'),
        ('Все страницы', 'Фирменный стиль', 'Логотип, цвета, шрифты и оформление соответствуют стилю компании и едины на всех страницах'),
        ('Все страницы', 'Тексты', 'Нет опечаток, ошибок и устаревших сведений; формулировки согласованы'),
        ('Все страницы', 'Изображения и видео', 'Изображения чёткие, соответствуют продуктам; видео воспроизводится'),
        ('Все страницы', 'Ссылки и кнопки', 'Нет «битых» ссылок; все кнопки выполняют ожидаемое действие'),
        (H + 'несуществующая-страница/', 'Страница 404', 'Показывается оформленная страница «не найдено» со ссылкой на главную'),
        ('Все страницы', 'Заголовки вкладок браузера', 'У каждой страницы понятный заголовок во вкладке'),
    ]),
    ('Главная', [
        (H, 'Первый экран', 'Понятно, чем занимается компания; есть кнопки заявки'),
        (H, 'Блок «Решения для практического обучения БПЛА»', 'Показаны все 5 продуктов, ссылки ведут на их страницы'),
        (H, 'Блок «Какая задача стоит перед вами?»', 'Варианты задач описаны верно, кнопки открывают нужные формы'),
        (H, 'Блок «Класс БПЛА под вашу образовательную программу»', 'Описание корректно, кнопка «Рассчитать класс» открывает форму'),
        (H, 'Блок «От задачи до занятий»', 'Этапы работы с клиентом описаны верно'),
        (H, 'Блок «Собственная разработка…»', 'Сведения о компании корректны'),
        (H, 'Блок «Документы для изучения и закупки»', 'Перечень документов корректен, кнопки работают'),
        (H, 'Блок «События и разработки»', 'Выводятся последние новости, ссылки открывают новость'),
    ]),
    ('Продукты', [
        (H + 'product/', 'Страница «Продукты»', 'Открывается каталог решений и оборудования'),
        (H + 'product/pchelka/', 'Учебный комплекс «Пчёлка»', 'Описание, состав, возможности и изображения корректны'),
        (H + 'product/constructor/', 'Дрон-конструктор', 'Описание, возможности и изображения корректны'),
        (H + 'product/lab/', 'Лабораторный стенд', 'Описание, возможности и изображения корректны'),
        (H + 'product/simulator/', 'АТС Симулятор', 'Описание корректно; указано, что симулятор — часть учебного комплекса «Пчёлка»'),
        (H + 'product/classroom/', 'Класс БПЛА', 'Описание корректно, кнопка «Рассчитать класс» открывает форму'),
        ('Страницы продуктов', 'Блок «Другие решения линейки»', 'Ссылки ведут на остальные продукты'),
        ('Страницы продуктов', 'Названия и характеристики', 'Названия продуктов и технические данные совпадают с официальными'),
    ]),
    ('Разделы', [
        (H + 'education/', 'Для образования', 'Программа обучения и темы курса описаны верно'),
        (H + 'procurement/', 'Как закупить', 'Сведения о закупке по 44-ФЗ, 223-ФЗ и другим способам корректны'),
        (H + 'documents/', 'Документы', 'Перечень документов (технические описания, спецификации, методические материалы) корректен, запрос документа открывает форму'),
        (H + 'about/', 'О компании', 'Сведения о компании корректны и актуальны'),
        (H + 'contacts/', 'Контакты', 'Телефон +7 (987) 499-74-97 и email aviatechnosoft@yandex.ru актуальны, ссылки срабатывают по нажатию'),
        (H + 'equipment/', 'Оборудование и компоненты', 'Перечень оборудования корректен'),
    ]),
    ('Новости', [
        (H + 'news/', 'Список новостей', 'Новости выводятся по дате, есть переход на следующие страницы'),
        (H + 'news/', 'Отдельная новость', 'Заголовок, дата, текст и изображения отображаются корректно'),
        ('Админка WordPress', 'Добавление новости', 'Сотрудник заказчика может сам создать, отредактировать и опубликовать новость'),
    ]),
    ('Формы заявок', [
        ('Кнопки «Получить КП»', 'Форма «Получить КП»', 'Форма открывается; есть выбор продукта и отметка «КП для обоснования цены контракта»; заявка отправляется'),
        ('Кнопки «Рассчитать класс»', 'Форма «Рассчитать класс»', 'Есть поля «Тип учреждения», «Число учащихся», «Срок запуска и бюджет»; заявка отправляется'),
        ('Кнопки «Запросить демонстрацию»', 'Форма «Запросить демонстрацию»', 'Есть выбор формата (онлайн / очно / обсудить); заявка отправляется'),
        ('Кнопки «Получить документацию»', 'Форма «Получить документацию»', 'Название документа подставляется автоматически, поле контакта принимает только email; заявка отправляется'),
        ('Все формы', 'Получение писем', 'Письмо по каждой заявке приходит на корпоративную почту и содержит все заполненные поля и тип заявки'),
        ('Все формы', 'Проверка заполнения', 'При пустых обязательных полях или без согласия на обработку данных показывается понятная подсказка'),
        ('Все формы', 'Сообщение после отправки', 'После отправки появляется сообщение об успехе; форма закрывается кнопкой × и клавишей Esc'),
    ]),
    ('Устройства и браузеры', [
        ('Все страницы', 'Телефон (ширина ~390 px)', 'Текст читается без увеличения, нет горизонтальной прокрутки, кнопки удобно нажимать'),
        ('Все страницы', 'Мобильное меню', 'Меню открывается и закрывается, все пункты доступны'),
        ('Все страницы', 'Планшет', 'Страницы отображаются корректно'),
        ('Все страницы', 'Компьютер (ширина ~1440 px)', 'Страницы отображаются корректно'),
        ('Все страницы', 'Браузеры', 'Сайт корректно работает в Яндекс Браузере, Chrome, Firefox, Safari и Edge'),
        ('Все страницы', 'Скорость загрузки', 'Страницы открываются быстро, в том числе через мобильный интернет'),
    ]),
    ('Юридические и служебные сведения', [
        ('Все формы', 'Согласие на обработку персональных данных', 'Есть отметка согласия и ссылка на утверждённый текст политики (сейчас текст не добавлен — ждёт заказчика)'),
        ('Все страницы', 'Политика обработки персональных данных', 'Опубликована страница с утверждённым текстом (152-ФЗ)'),
        ('Все страницы', 'Реквизиты компании', 'Реквизиты указаны и верны (после предоставления заказчиком)'),
        ('Все страницы', 'Защищённое соединение', 'Сайт открывается по https, браузер не показывает предупреждений'),
    ]),
]

wb = Workbook()

# --- Инструкция ---
ws = wb.active
ws.title = 'Инструкция'
ws.column_dimensions['A'].width = 100
lines = [
    ('Лист приёмки сайта aviatechnosoft.ru', title),
    ('ООО «АвиаТехноСофт», корпоративный сайт', bold),
    ('', font),
    ('Как заполнять', bold),
    ('1. Откройте лист «Лист приёмки» и пройдите пункты по порядку.', font),
    ('2. Заполняйте только ячейки с жёлтой заливкой.', font),
    ('3. В колонке «Оценка» выберите значение из списка:', font),
    ('     • Принято — всё соответствует ожидаемому результату;', font),
    ('     • Замечание — есть недочёт, который нужно исправить, но сайт работоспособен;', font),
    ('     • Критично — ошибка мешает запуску сайта;', font),
    ('     • Не проверялось — пункт пропущен.', font),
    ('4. Для «Замечание» и «Критично» опишите проблему в «Комментарии» и укажите устройство и браузер.', font),
    ('5. Замечания, которых нет в списке, запишите на лист «Доп. замечания» (там есть пример заполнения).', font),
    ('6. На листе «Данные от заказчика» отметьте, какие материалы вы можете предоставить.', font),
    ('7. Лист «Итог» считается автоматически; внизу укажите решение, ФИО и дату.', font),
    ('', font),
    ('Скриншоты можно прикладывать отдельными файлами, указав в комментарии их имя.', font),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = f
ws.cell(row=19, column=1, value='Ячейки для заполнения').font = bold
ws.cell(row=20, column=1).fill = inputfill
ws.cell(row=20, column=1, value='— так выделены ячейки, которые заполняет заказчик').font = font

# --- Лист приёмки ---
ws = wb.create_sheet('Лист приёмки')
headers = ['№', 'Раздел', 'Страница / где проверять', 'Что проверить', 'Ожидаемый результат',
           'Оценка', 'Комментарий заказчика', 'Устройство / браузер']
widths = [6, 18, 34, 34, 55, 16, 45, 22]
ws.append(['Проверил (ФИО):', '', '', 'Дата проверки:', ''])
for col in (2, 5):
    ws.cell(row=1, column=col).fill = inputfill
ws.merge_cells('B1:C1')
for col in (1, 4):
    ws.cell(row=1, column=col).font = bold
ws.append([])
ws.append(headers)
HR = 3
for i, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=HR, column=i)
    c.font, c.fill, c.alignment, c.border = hfont, hfill, Alignment(wrap_text=True, vertical='center'), border
    ws.column_dimensions[c.column_letter].width = w
r = HR + 1
n = 0
first = r
for section, rows in items:
    ws.cell(row=r, column=1, value=section)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    for col in range(1, 9):
        ws.cell(row=r, column=col).fill = secfill
        ws.cell(row=r, column=col).border = border
    ws.cell(row=r, column=1).font = bold
    r += 1
    for page, what, expected in rows:
        n += 1
        vals = [n, section, page, what, expected, None, None, None]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.font, c.alignment, c.border = font, wrap, border
            if col >= 6:
                c.fill = inputfill
        r += 1
last = r - 1
ws.freeze_panes = ws.cell(row=HR + 1, column=1)
dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=True)
dv.error = 'Выберите значение из списка'
ws.add_data_validation(dv)
dv.add(f'F{first}:F{last}')
for status, color in (('Принято', 'C6EFCE'), ('Замечание', 'FFEB9C'), ('Критично', 'FFC7CE')):
    ws.conditional_formatting.add(
        f'F{first}:F{last}',
        FormulaRule(formula=[f'$F{first}="{status}"'], fill=PatternFill('solid', fgColor=color)))
ws.auto_filter.ref = f'A{HR}:H{last}'
ws.print_title_rows = f'{HR}:{HR}'
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
RANGE = f"'Лист приёмки'!$F${first}:$F${last}"
NUMS = f"'Лист приёмки'!$A${first}:$A${last}"

# --- Доп. замечания ---
ws = wb.create_sheet('Доп. замечания')
h2 = ['№', 'Страница (адрес)', 'Описание проблемы или пожелания', 'Важность', 'Устройство / браузер', 'Скриншот (имя файла)']
w2 = [6, 36, 60, 16, 22, 26]
ws.append(h2)
for i, w in enumerate(w2, 1):
    c = ws.cell(row=1, column=i)
    c.font, c.fill, c.alignment, c.border = hfont, hfill, Alignment(wrap_text=True, vertical='center'), border
    ws.column_dimensions[c.column_letter].width = w
example = ['Пример', H + 'contacts/', 'Хотим добавить карту проезда к офису', 'Замечание', 'iPhone 13, Safari', 'contacts-map.png']
for col, v in enumerate(example, 1):
    c = ws.cell(row=2, column=col, value=v)
    c.font, c.alignment, c.border = Font(name=F, size=10, italic=True, color='808080'), wrap, border
for row in range(3, 33):
    ws.cell(row=row, column=1, value=row - 2).font = font
    for col in range(1, 7):
        c = ws.cell(row=row, column=col)
        c.border, c.alignment = border, wrap
        c.font = font
        if col > 1:
            c.fill = inputfill
dv2 = DataValidation(type='list', formula1='"Замечание,Критично,Пожелание"', allow_blank=True)
ws.add_data_validation(dv2)
dv2.add('D3:D32')
ws.freeze_panes = 'A2'

# --- Данные от заказчика ---
ws = wb.create_sheet('Данные от заказчика')
h3 = ['№', 'Материал', 'Где используется', 'Статус', 'Комментарий заказчика']
w3 = [6, 40, 45, 18, 45]
ws.append(h3)
for i, w in enumerate(w3, 1):
    c = ws.cell(row=1, column=i)
    c.font, c.fill, c.alignment, c.border = hfont, hfill, Alignment(wrap_text=True, vertical='center'), border
    ws.column_dimensions[c.column_letter].width = w
needs = [
    ('Текст политики обработки персональных данных', 'Формы заявок, отдельная страница политики'),
    ('Реквизиты компании (ИНН, ОГРН, юридический адрес)', 'Подвал, страница «О компании», «Контакты»'),
    ('Цены или порядок их запроса', 'Страницы продуктов, «Как закупить»'),
    ('Составы комплектов продуктов', 'Страницы продуктов'),
    ('Сведения о включении в реестры (реестр ПО, реестр российской продукции и т. п.)', 'Страницы продуктов, «Как закупить», «Документы»'),
    ('Документы для скачивания или отправки по запросу', 'Страница «Документы»'),
    ('Кейсы внедрения (не менее двух согласованных)', 'Раздел кейсов (скрыт до появления материалов)'),
    ('Email для получения заявок с сайта', 'Настройка отправки форм'),
    ('Фотографии и видео продуктов (при необходимости замены)', 'Главная, страницы продуктов'),
]
dv3 = DataValidation(type='list', formula1='"Предоставим,Уже передано,Не требуется"', allow_blank=True)
ws.add_data_validation(dv3)
for i, (m, where) in enumerate(needs, 1):
    for col, v in enumerate([i, m, where, None, None], 1):
        c = ws.cell(row=i + 1, column=col, value=v)
        c.font, c.alignment, c.border = font, wrap, border
        if col >= 4:
            c.fill = inputfill
dv3.add(f'D2:D{len(needs) + 1}')
ws.freeze_panes = 'A2'

# --- Итог ---
ws = wb.create_sheet('Итог')
ws.column_dimensions['A'].width = 34
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 14
ws['A1'] = 'Итог приёмки сайта'
ws['A1'].font = title
for col, h in enumerate(['Оценка', 'Пунктов', 'Доля'], 1):
    c = ws.cell(row=3, column=col, value=h)
    c.font, c.fill, c.border = hfont, hfill, border
for i, s in enumerate(STATUSES):
    row = 4 + i
    ws.cell(row=row, column=1, value=s)
    ws.cell(row=row, column=2, value=f'=COUNTIF({RANGE},A{row})')
    ws.cell(row=row, column=3, value=f'=IF($B$9=0,0,B{row}/$B$9)')
row = 8
ws.cell(row=row, column=1, value='Не заполнено')
ws.cell(row=row, column=2, value='=B9-SUM(B4:B7)')
ws.cell(row=row, column=3, value=f'=IF($B$9=0,0,B{row}/$B$9)')
ws.cell(row=9, column=1, value='Всего пунктов')
ws.cell(row=9, column=2, value=f'=COUNT({NUMS})')
ws.cell(row=9, column=3, value='=SUM(C4:C8)')
for rr in range(4, 10):
    for col in range(1, 4):
        c = ws.cell(row=rr, column=col)
        c.border = border
        c.font = bold if rr == 9 else font
        if col == 3:
            c.number_format = '0%'
ws['A10'] = 'Считается автоматически по листу «Лист приёмки».'
ws['A10'].font = Font(name=F, size=9, italic=True, color='808080')
ws['A11'] = 'Доп. замечаний записано'
ws['B11'] = "=COUNTA('Доп. замечания'!C3:C32)"
for c in (ws['A11'], ws['B11']):
    c.font, c.border = font, border

ws['A13'] = 'Решение заказчика'
ws['A13'].font = bold
ws['B13'].fill = inputfill
ws.merge_cells('B13:E13')
dv4 = DataValidation(type='list', formula1='"Принять,Принять с замечаниями,Доработать"', allow_blank=True)
ws.add_data_validation(dv4)
dv4.add('B13')
for rr, label in ((14, 'ФИО, должность'), (15, 'Дата'), (16, 'Подпись')):
    ws.cell(row=rr, column=1, value=label).font = bold
    ws.cell(row=rr, column=2).fill = inputfill
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=5)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('items', n)
